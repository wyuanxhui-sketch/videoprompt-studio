"""All credentials and provider responses are synthetic; never call an API."""
import json
import os
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch
from urllib.error import HTTPError, URLError

import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from streamlit.proto.TextInput_pb2 import TextInput

from ai_storyboard import (AIError, AI_EXAMPLE, MAX_INPUT, MAX_SHOTS, MAX_TOKENS,
                          MODELS, TIMEOUT, build_payload, claim_request, generate_ai,
                          parse_storyboard, resolve_api_key, _NoRedirect)
from generator import export_csv, export_excel
from storyboard import COLUMNS

APP = str(Path(__file__).resolve().parents[1] / "app.py")
FAKE_KEY = "mock-only-not-a-real-key"


def fixture_shots():
    # Demonstrates the contract, not evidence of a live model's narrative quality.
    events = [
        ("佛光寺大殿正立面", "古装少女从画面左侧向右前方走近正立面，以全景建立人物与大殿的空间关系", "全景", "古装", "第四幕"),
        ("佛光寺大殿", "沿少女视线观察斗拱；以可靠图纸辅助表达双槽空间关系，待参考资料核对", "近景", "古装", "第四幕"),
        ("古建至展陈过渡空间", "同一少女继续向右前方行走，脚步动作匹配，古建背景逐渐转入现代展陈", "中景", "古装过渡现代服饰", "第五幕"),
        ("现代展陈空间", "沿用上一镜头步态与机位，衣服由古装过渡至现代服饰，脸部和发型保持一致", "中景", "现代服饰", "第五幕"),
        ("现代展陈空间", "现代服饰少女向右前方走近悬挂的斗拱状展陈结构，以宽阔空间交代观看距离", "全景", "现代服饰", "第六幕"),
    ]
    rows = []
    for i, (place, content, size, costume, act) in enumerate(events, 1):
        note = (f"AI视频生成提示词：{act}，电影写实，自然均匀光线，{size}，平视跟拍；{content}；{costume}。"
                "\n音效：连续脚步与环境声。\n转场：运动匹配与渐变。"
                "\n连续性约束：同一少女脸部、发型不变，向右前方行走；双槽与斗拱关系依据实拍图纸核对，不编造结构尺寸。"
                "\n后期制作：复杂变装分段生成，匹配步态与构图后合成拼接。")
        rows.append(dict(zip(COLUMNS, [f"{i:02d}", place, content, size, "平视", "跟拍", 7, "无", note, ""])))
    return rows


def response_body(content=None, finish="stop", usage=None):
    return {"choices": [{"finish_reason": finish, "message": {"content": content if content is not None else json.dumps({"shots": fixture_shots()}, ensure_ascii=False)}}],
            "usage": usage if usage is not None else {"prompt_tokens": 120, "completion_tokens": 800, "total_tokens": 920}}


def transport_for(body):
    raw = json.dumps(body, ensure_ascii=False).encode()
    response = Mock()
    response.__enter__ = Mock(return_value=BytesIO(raw))
    response.__exit__ = Mock(return_value=False)
    return Mock(return_value=response)


class AIClientTests(unittest.TestCase):
    def setUp(self):
        self.payload = build_payload(AI_EXAMPLE, "电影写实", 7, next(iter(MODELS)), 35, "同一少女", 12)

    def test_normal_json_and_usage(self):
        transport = transport_for(response_body())
        frame, usage = generate_ai(self.payload, FAKE_KEY, transport)
        self.assertEqual(list(frame.columns), COLUMNS)
        self.assertEqual(len(frame), 5)
        self.assertEqual(usage["total_tokens"], 920)
        request = transport.call_args.args[0]
        self.assertEqual(request.full_url, "https://api.deepseek.com/chat/completions")
        self.assertEqual(request.get_header("Authorization"), f"Bearer {FAKE_KEY}")
        self.assertNotIn(FAKE_KEY, request.data.decode())
        self.assertEqual(transport.call_args.kwargs["timeout"], TIMEOUT)
        self.assertEqual(transport.call_count, 1)

    def test_json_error(self):
        with self.assertRaisesRegex(AIError, "有效 JSON") as caught:
            generate_ai(self.payload, FAKE_KEY, transport_for(response_body("{bad")))
        self.assertEqual(caught.exception.usage["total_tokens"], 920)

    def test_missing_field(self):
        rows = fixture_shots()
        del rows[0]["地点"]
        with self.assertRaisesRegex(AIError, "字段缺失"):
            parse_storyboard(json.dumps({"shots": rows}), 12)

    def test_empty_response(self):
        for body in (response_body(""), {"choices": []}, None, [], response_body(" ")):
            with self.subTest(body=body), self.assertRaises(AIError):
                generate_ai(self.payload, FAKE_KEY, transport_for(body))

    def test_timeout_no_retry(self):
        for timeout in (TimeoutError(FAKE_KEY), URLError(TimeoutError(FAKE_KEY))):
            transport = Mock(side_effect=timeout)
            with self.assertRaisesRegex(AIError, "超时") as caught:
                generate_ai(self.payload, FAKE_KEY, transport)
            self.assertNotIn(FAKE_KEY, str(caught.exception))
            self.assertEqual(transport.call_count, 1)

    def test_no_key_never_calls(self):
        transport = Mock()
        with self.assertRaisesRegex(AIError, "API Key"):
            generate_ai(self.payload, "", transport)
        transport.assert_not_called()

    def test_network_and_http_errors_safe_no_retry(self):
        for error in [URLError(FAKE_KEY)] + [HTTPError("https://api.deepseek.com", code, FAKE_KEY, {}, None) for code in (401, 402, 429, 500, 302)]:
            transport = Mock(side_effect=error)
            with self.subTest(error=type(error)), self.assertRaises(AIError) as caught:
                generate_ai(self.payload, FAKE_KEY, transport)
            self.assertNotIn(FAKE_KEY, str(caught.exception))
            self.assertEqual(transport.call_count, 1)

    def test_truncation(self):
        with self.assertRaisesRegex(AIError, "截断") as caught:
            generate_ai(self.payload, FAKE_KEY, transport_for(response_body(finish="length")))
        self.assertEqual(caught.exception.usage["total_tokens"], 920)

    def test_invalid_response_envelope(self):
        for body in ({}, {"choices": "bad"}, {"choices": [None]}, {"choices": [{"message": {}}]}):
            with self.subTest(body=body), self.assertRaises(AIError):
                generate_ai(self.payload, FAKE_KEY, transport_for(body))

    def test_invalid_schema_values(self):
        for column, value in [("时长（秒）", 0), ("时长（秒）", True), ("时长（秒）", float("nan")), ("时长（秒）", 61),
                              ("地点", None), ("画面内容", ""), ("备注", "只有提示词"), ("镜头号", "02"), ("配图", "unexpected")]:
            rows = fixture_shots()
            rows[0][column] = value
            with self.subTest(column=column, value=value), self.assertRaises(AIError):
                parse_storyboard(json.dumps({"shots": rows}), 12)

    def test_arrays_and_duplicate_keys_rejected(self):
        for content in ('{"shots":[],"shots":[]}', '{"shots":[]}', '[]', '{"shots":{},"extra":1}', json.dumps({"shots": fixture_shots()})):
            with self.subTest(content=content[:40]), self.assertRaises(AIError):
                parse_storyboard(content, 1)

    def test_limits_and_prompt_contract(self):
        self.assertEqual(self.payload["max_tokens"], MAX_TOKENS)
        self.assertEqual(self.payload["response_format"], {"type": "json_object"})
        self.assertEqual(self.payload["thinking"], {"type": "disabled"})
        data = json.loads(self.payload["messages"][1]["content"])
        self.assertEqual(data["文案"], AI_EXAMPLE)
        self.assertEqual(data["目标总时长（秒）"], 35)
        self.assertEqual(list(data["JSON格式示例"]["shots"][0]), COLUMNS)
        for concept in ("不按标点", "同一人物", "双槽", "不得编造", "分段生成", "服装变化"):
            self.assertIn(concept, self.payload["messages"][0]["content"])
        for kwargs in ({"source": ""}, {"source": "字" * (MAX_INPUT + 1)}, {"continuity": "字" * 1001},
                       {"max_shots": MAX_SHOTS + 1}, {"max_shots": True}, {"model": "unverified"}, {"target": float("nan")}, {"target": 601}):
            args = dict(source=AI_EXAMPLE, style="通用", duration=7, model=next(iter(MODELS)))
            args.update(kwargs)
            with self.subTest(kwargs=list(kwargs)), self.assertRaises(AIError):
                build_payload(**args)

    def test_public_mode_does_not_use_developer_key(self):
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "mock-developer", "VIDEOPROMPT_LOCAL_DEV": "0"}):
            self.assertEqual(resolve_api_key(""), "")
            self.assertEqual(resolve_api_key("", True), "")
            self.assertEqual(resolve_api_key("mock-user"), "mock-user")
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "mock-developer", "VIDEOPROMPT_LOCAL_DEV": "1"}):
            self.assertEqual(resolve_api_key(""), "")
            self.assertEqual(resolve_api_key("", True), "")
            self.assertEqual(resolve_api_key("mock-user", True), "mock-user")

    def test_session_isolation_and_duplicate_claim(self):
        first, second = {}, {}
        self.assertTrue(claim_request(first, self.payload))
        self.assertFalse(claim_request(first, self.payload))
        first["ai_busy"] = False
        self.assertFalse(claim_request(first, self.payload))
        self.assertTrue(claim_request(second, self.payload))
        first.pop("ai_attempt")
        self.assertTrue(claim_request(first, self.payload))
        self.assertNotIn(FAKE_KEY, repr(first))

    def test_key_echo_and_redirect_rejected(self):
        with self.assertRaisesRegex(AIError, "敏感"):
            generate_ai(self.payload, FAKE_KEY, transport_for(response_body(FAKE_KEY)))
        self.assertIsNone(_NoRedirect().redirect_request(None, None, 302, "", {}, "https://other.invalid"))

    def test_usage_whitelist(self):
        _, usage = generate_ai(self.payload, FAKE_KEY, transport_for(response_body(usage={"secret": FAKE_KEY, "total_tokens": 9, "prompt_tokens": FAKE_KEY})))
        self.assertEqual(usage, {"total_tokens": 9})

    def test_response_length_limit(self):
        response = Mock()
        response.__enter__ = Mock(return_value=BytesIO(b" " * 512001))
        response.__exit__ = Mock(return_value=False)
        with self.assertRaisesRegex(AIError, "长度限制"):
            generate_ai(self.payload, FAKE_KEY, Mock(return_value=response))

    def test_edit_add_delete_and_exports(self):
        frame, _ = generate_ai(self.payload, FAKE_KEY, transport_for(response_body()))
        frame.loc[0, "备注"] = "=人工修改"
        frame.loc[len(frame)] = frame.iloc[0].copy()
        frame = frame.drop(index=1)
        csv = pd.read_csv(BytesIO(export_csv(frame)), dtype={"镜头号": str})
        self.assertEqual(list(csv.columns), COLUMNS)
        self.assertEqual(csv.iloc[0]["备注"], "=人工修改")
        sheet = load_workbook(BytesIO(export_excel(frame))).active
        self.assertEqual([c.value for c in sheet[1]], COLUMNS)
        self.assertEqual(sheet.max_row, len(frame) + 1)
        self.assertEqual(sheet["I2"].value, "=人工修改")
        self.assertEqual(sheet["I2"].data_type, "s")
        self.assertNotIn(FAKE_KEY.encode(), export_csv(frame))


class AIWebTests(unittest.TestCase):
    def setUp(self):
        # Default network boundary blocked for every web test.
        self.network = patch("ai_storyboard.build_opener")
        self.opener = self.network.start()
        self.opener.return_value.open.side_effect = AssertionError("Real API forbidden")
        self.addCleanup(self.network.stop)
        self.env = patch.dict(os.environ, {"VIDEOPROMPT_LOCAL_DEV": "0", "DEEPSEEK_API_KEY": "mock-developer"})
        self.env.start()
        self.addCleanup(self.env.stop)

    def make_app(self):
        app = AppTest.from_file(APP).run()
        app.radio[0].set_value("AI智能模式").run()
        app.text_area[0].set_value(AI_EXAMPLE).run()
        return app

    @staticmethod
    def button(app, label):
        return next(b for b in app.button if b.label == label)

    def test_no_call_on_refresh_or_missing_key(self):
        app = self.make_app()
        app.run()
        self.button(app, "开始AI生成").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("API Key" in error.value for error in app.error))
        self.opener.assert_not_called()
        self.assertEqual(app.text_input[0].proto.type, TextInput.PASSWORD)

    def test_paid_button_only_and_duplicate_guard(self):
        frame = parse_storyboard(json.dumps({"shots": fixture_shots()}), 12)
        with patch("ai_storyboard.generate_ai", return_value=(frame, {"total_tokens": 920})) as api:
            app = self.make_app()
            app.text_input[0].set_value(FAKE_KEY).run()
            api.assert_not_called()
            self.button(app, "开始AI生成").click().run()
            self.assertFalse(app.exception)
            self.assertEqual(api.call_count, 1)
            self.assertIn("成功", app.session_state["ai_status"])
            app.run()
            self.button(app, "开始AI生成").click().run()
            self.assertEqual(api.call_count, 1)
            self.button(app, "解除相同请求保护（下次开始AI生成可能再次计费）").click().run()
            self.assertEqual(api.call_count, 1)
            self.button(app, "开始AI生成").click().run()
            self.assertEqual(api.call_count, 2)
            self.assertNotIn(FAKE_KEY, app.code[-1].value)

    def test_failure_retains_previous_result(self):
        app = AppTest.from_file(APP).run()
        app.button[0].click().run()
        app.button[1].click().run()
        old = app.dataframe[0].value.copy()
        app.radio[0].set_value("AI智能模式").run()
        app.text_input[0].set_value(FAKE_KEY).run()
        with patch("ai_storyboard.generate_ai", side_effect=AIError("API 请求超时")) as api:
            self.button(app, "开始AI生成").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(api.call_count, 1)
        pd.testing.assert_frame_equal(old, app.dataframe[0].value)
        self.assertIn("失败", app.session_state["ai_status"])
        self.assertFalse(app.session_state["ai_busy"])

    def test_two_users_and_new_session_keys(self):
        frame = parse_storyboard(json.dumps({"shots": fixture_shots()}), 12)
        first, second = self.make_app(), self.make_app()
        first.text_input[0].set_value("mock-user-a").run()
        self.assertEqual(second.text_input[0].value, "")
        second.text_input[0].set_value("mock-user-b").run()
        with patch("ai_storyboard.generate_ai", return_value=(frame, {})) as api:
            self.button(first, "开始AI生成").click().run()
            self.button(second, "开始AI生成").click().run()
        self.assertEqual([c.args[1] for c in api.call_args_list], ["mock-user-a", "mock-user-b"])
        self.assertEqual(self.make_app().text_input[0].value, "")

    def test_editor_changes_do_not_call_api(self):
        frame = parse_storyboard(json.dumps({"shots": fixture_shots()}), 12)
        with patch("ai_storyboard.generate_ai", return_value=(frame, {})) as api:
            app = self.make_app()
            app.text_input[0].set_value(FAKE_KEY).run()
            self.button(app, "开始AI生成").click().run()
            revision = app.session_state["revision"]
            app.session_state[f"editor_{revision}"] = {"edited_rows": {0: {"备注": "人工修改"}}, "added_rows": [], "deleted_rows": [1]}
            app.run()
            self.assertFalse(app.exception)
            self.assertIn("人工修改", app.code[-1].value)
            self.assertEqual(api.call_count, 1)


if __name__ == "__main__":
    unittest.main()

